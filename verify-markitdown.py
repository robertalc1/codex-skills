"""Offline smoke test for the isolated MarkItDown installation."""
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from markitdown import MarkItDown
from openpyxl import Workbook
from pptx import Presentation


with TemporaryDirectory(prefix="codex-markitdown-") as directory:
    root = Path(directory)
    (root / "sample.html").write_text("<h1>Collection test</h1><p>offline sentinel</p>", encoding="utf-8")
    with ZipFile(root / "sample.docx", "w") as doc:
        doc.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        doc.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        doc.writestr("word/document.xml", '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>offline sentinel</w:t></w:r></w:p></w:body></w:document>')
    book = Workbook()
    book.active.append(["Name", "Value"])
    book.active.append(["offline sentinel", 159])
    book.save(root / "sample.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Collection test"
    slide.placeholders[1].text = "offline sentinel"
    slides.save(root / "sample.pptx")
    converter = MarkItDown(enable_plugins=False)
    for extension in ("html", "docx", "xlsx", "pptx"):
        result = converter.convert_local(str(root / f"sample.{extension}"))
        if "offline sentinel" not in result.text_content:
            raise AssertionError(f"Missing expected content in {extension} conversion")
        print(f"{extension}: passed")
